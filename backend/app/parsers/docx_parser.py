"""
Universal DOCX Document Parser with deep segment extraction.
Extracts body paragraphs, tables, headers/footers, comments, tracked changes,
hidden text runs (w:vanish), footnotes, image alt text, and metadata properties.
Enforces decompression bomb checks before parsing.
"""
import io
import hashlib
import zipfile
import re
import uuid
from defusedxml import ElementTree as ET
from typing import List, Dict, Any
from .base import BaseParser, ExtractedContent
from .sniff import check_zip_bomb
from app.models.schemas import Segment

try:
    import docx
    _HAS_DOCX = True
except ImportError:
    docx = None
    _HAS_DOCX = False

W_NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}

class DOCXParser(BaseParser):
    supported_types = ['application/vnd.openxmlformats-officedocument.wordprocessingml.document']
    MAX_EMBEDDED_IMAGES = 25

    def can_parse(self, mime_type: str, filename: str) -> bool:
        return mime_type in self.supported_types or filename.lower().endswith(('.docx', '.doc'))

    async def parse(self, content: bytes, filename: str = '', mime_type: str = '') -> ExtractedContent:
        extracted = ExtractedContent(
            text='',
            source_type='docx'
        )
        extracted.content_hash = hashlib.sha256(content).hexdigest()
        segments: List[Segment] = []
        embedded_image_texts = []

        # 1. Decompression Bomb Check
        is_safe, bomb_reason = check_zip_bomb(content)
        if not is_safe:
            extracted.suspicious_elements.append({
                "type": "zip_bomb_detected",
                "description": bomb_reason
            })
            extracted.extraction_warnings.append(f"File blocked: {bomb_reason}")
            return extracted

        # 2. Inspect ZIP structure for macros, external rels, comments, hidden text
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as docx_zip:
                namelist = docx_zip.namelist()

                # Word documents can carry prompt text as pixels while their
                # visible paragraphs look benign. Inspect embedded raster media
                # with the same bounded OCR pipeline as uploaded images.
                media_names = [
                    name for name in namelist
                    if name.startswith("word/media/") and not name.endswith("/")
                ]
                if len(media_names) > self.MAX_EMBEDDED_IMAGES:
                    extracted.extraction_warnings.append(
                        f"DOCX embedded-image budget exceeded ({self.MAX_EMBEDDED_IMAGES})"
                    )
                from .image_parser import ImageParser
                image_parser = ImageParser()
                for media_name in media_names[:self.MAX_EMBEDDED_IMAGES]:
                    try:
                        image_bytes = docx_zip.read(media_name)
                        image_result = await image_parser.parse(image_bytes, media_name)
                        extracted.extraction_warnings.extend(image_result.extraction_warnings)
                        for image_segment in image_result.segments:
                            image_segment.source_type = "docx"
                            image_segment.location = (
                                f"docx:embedded_image:{media_name}:{image_segment.location}"
                            )
                            segments.append(image_segment)
                        if image_result.text.strip():
                            embedded_image_texts.append(
                                f"[embedded image {media_name} OCR]\n{image_result.text}"
                            )
                    except Exception as exc:
                        extracted.extraction_warnings.append(
                            f"DOCX embedded image OCR failed: {type(exc).__name__}"
                        )

                # Macros check
                if 'word/vbaProject.bin' in namelist:
                    extracted.suspicious_elements.append({
                        "type": "macro",
                        "description": "VBA Macro found in DOCX."
                    })

                # External relationships check
                for item in namelist:
                    if item.endswith('.rels'):
                        rels_data = docx_zip.read(item).decode('utf-8', errors='ignore')
                        urls = re.findall(r'Target="([^"]+)"', rels_data)
                        for url in urls:
                            if url.startswith('http'):
                                extracted.links.append(url)
                            if "external" in rels_data.lower() and (url.startswith('http') or url.startswith('\\\\')):
                                extracted.suspicious_elements.append({
                                    "type": "external_relationship",
                                    "description": f"External relationship found: {url}"
                                })

                # Comments XML extraction (hidden comments)
                if 'word/comments.xml' in namelist:
                    try:
                        comments_xml = docx_zip.read('word/comments.xml')
                        tree = ET.fromstring(comments_xml)
                        for i, comment_elem in enumerate(tree.iterfind('.//w:comment', W_NS)):
                            c_text = "".join(comment_elem.itertext()).strip()
                            if c_text:
                                extracted.comments.append(c_text)
                                segments.append(Segment(
                                    id=f"seg-docx-comment-{i}-{uuid.uuid4().hex[:6]}",
                                    text=c_text,
                                    source_type="docx",
                                    location=f"docx:comment:{i+1}",
                                    visibility="hidden",
                                    trust="UNTRUSTED"
                                ))
                    except Exception:
                        extracted.extraction_warnings.append("DOCX comment extraction failed")

                # Footnotes / Endnotes XML extraction
                for fn_file, loc_name in [('word/footnotes.xml', 'footnote'), ('word/endnotes.xml', 'endnote')]:
                    if fn_file in namelist:
                        try:
                            note_xml = docx_zip.read(fn_file)
                            tree = ET.fromstring(note_xml)
                            for i, note_elem in enumerate(tree.iterfind(f'.//w:{loc_name}', W_NS)):
                                n_text = "".join(note_elem.itertext()).strip()
                                if n_text:
                                    segments.append(Segment(
                                        id=f"seg-docx-{loc_name}-{i}-{uuid.uuid4().hex[:6]}",
                                        text=n_text,
                                        source_type="docx",
                                        location=f"docx:{loc_name}:{i+1}",
                                        visibility="visible",
                                        trust="UNTRUSTED"
                                    ))
                        except Exception:
                            extracted.extraction_warnings.append("DOCX subdocument extraction failed")

                # Document.xml hidden text runs (w:vanish) & tracked changes (w:ins / w:del)
                if 'word/document.xml' in namelist:
                    try:
                        doc_xml = docx_zip.read('word/document.xml')
                        tree = ET.fromstring(doc_xml)

                        # Hidden text runs (w:vanish)
                        for i, r_elem in enumerate(tree.iterfind('.//w:r', W_NS)):
                            rPr = r_elem.find('w:rPr', W_NS)
                            if rPr is not None and rPr.find('w:vanish', W_NS) is not None:
                                vanish_text = "".join(r_elem.itertext()).strip()
                                if vanish_text:
                                    extracted.hidden_content.append(vanish_text)
                                    segments.append(Segment(
                                        id=f"seg-docx-hidden-run-{i}-{uuid.uuid4().hex[:6]}",
                                        text=vanish_text,
                                        source_type="docx",
                                        location=f"docx:hidden_run:{i+1}",
                                        visibility="hidden",
                                        trust="UNTRUSTED"
                                    ))

                        # Tracked changes (w:ins and w:delText)
                        for i, ins_elem in enumerate(tree.iterfind('.//w:ins', W_NS)):
                            ins_text = "".join(ins_elem.itertext()).strip()
                            if ins_text:
                                segments.append(Segment(
                                    id=f"seg-docx-tracked-ins-{i}-{uuid.uuid4().hex[:6]}",
                                    text=ins_text,
                                    source_type="docx",
                                    location=f"docx:tracked_insertion:{i+1}",
                                    visibility="hidden",
                                    trust="UNTRUSTED"
                                ))

                        for i, del_elem in enumerate(tree.iterfind('.//w:delText', W_NS)):
                            del_text = "".join(del_elem.itertext()).strip()
                            if del_text:
                                segments.append(Segment(
                                    id=f"seg-docx-tracked-del-{i}-{uuid.uuid4().hex[:6]}",
                                    text=del_text,
                                    source_type="docx",
                                    location=f"docx:tracked_deletion:{i+1}",
                                    visibility="hidden",
                                    trust="UNTRUSTED"
                                ))

                        # Image alt text (descr attribute in drawing docPr)
                        for i, docPr in enumerate(tree.iterfind('.//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr')):
                            descr = docPr.get('descr', '').strip()
                            if descr:
                                segments.append(Segment(
                                    id=f"seg-docx-alt-{i}-{uuid.uuid4().hex[:6]}",
                                    text=descr,
                                    source_type="docx",
                                    location=f"docx:alt_text:{i+1}",
                                    visibility="metadata",
                                    trust="UNTRUSTED"
                                ))
                    except Exception:
                        extracted.extraction_warnings.append("DOCX hidden content extraction failed")

                # Headers & Footers
                for hf_file in namelist:
                    if hf_file.startswith('word/header') or hf_file.startswith('word/footer'):
                        try:
                            hf_xml = docx_zip.read(hf_file)
                            tree = ET.fromstring(hf_xml)
                            hf_text = "".join(tree.itertext()).strip()
                            if hf_text:
                                hf_type = "header" if "header" in hf_file else "footer"
                                segments.append(Segment(
                                    id=f"seg-docx-{hf_type}-{uuid.uuid4().hex[:6]}",
                                    text=hf_text,
                                    source_type="docx",
                                    location=f"docx:{hf_type}",
                                    visibility="metadata",
                                    trust="UNTRUSTED"
                                ))
                        except Exception:
                            extracted.extraction_warnings.append("DOCX subdocument extraction failed")

        except zipfile.BadZipFile:
            extracted.extraction_warnings.append("Malformed or corrupted DOCX zip archive.")
            return extracted
        except Exception as e:
            extracted.extraction_warnings.append(f"Failed to inspect DOCX zip: {str(e)}")

        # 3. Text and Tables via python-docx
        if _HAS_DOCX:
            try:
                doc_stream = io.BytesIO(content)
                doc = docx.Document(doc_stream)
                text_parts = []

                for i, para in enumerate(doc.paragraphs):
                    p_text = para.text.strip()
                    if p_text:
                        text_parts.append(p_text)
                        segments.append(Segment(
                            id=f"seg-docx-p-{i}-{uuid.uuid4().hex[:6]}",
                            text=p_text,
                            source_type="docx",
                            location=f"docx:paragraph:{i+1}",
                            visibility="visible",
                            trust="UNTRUSTED"
                        ))

                for t_idx, table in enumerate(doc.tables):
                    table_rows = []
                    for row in table.rows:
                        row_vals = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_vals:
                            table_rows.append(" | ".join(row_vals))
                    if table_rows:
                        t_content = "\n".join(table_rows)
                        text_parts.append(t_content)
                        segments.append(Segment(
                            id=f"seg-docx-table-{t_idx}-{uuid.uuid4().hex[:6]}",
                            text=t_content,
                            source_type="docx",
                            location=f"docx:table:{t_idx+1}",
                            visibility="visible",
                            trust="UNTRUSTED"
                        ))

                extracted.text = "\n\n".join(text_parts)
                if embedded_image_texts:
                    extracted.text += "\n\n" + "\n\n".join(embedded_image_texts)

                # Core properties / Metadata
                core_props = doc.core_properties
                props = {
                    "author": core_props.author,
                    "title": core_props.title,
                    "subject": core_props.subject,
                    "keywords": core_props.keywords,
                    "comments": core_props.comments,
                }
                for prop_name, prop_val in props.items():
                    if prop_val and isinstance(prop_val, str) and prop_val.strip():
                        extracted.metadata[prop_name] = prop_val.strip()
                        segments.append(Segment(
                            id=f"seg-docx-meta-{prop_name}-{uuid.uuid4().hex[:6]}",
                            text=prop_val.strip(),
                            source_type="docx",
                            location=f"docx:metadata:{prop_name}",
                            visibility="metadata",
                            trust="UNTRUSTED"
                        ))

            except Exception as e:
                extracted.extraction_warnings.append(f"DOCX python-docx error: {str(e)}")
        else:
            extracted.extraction_warnings.append("python-docx is not installed.")

        extracted.segments = segments
        return extracted
